import io
import re
import urllib.parse
import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.styles import Font

# 🚀 행정안전부 도로명주소(주소정보누리집) 오픈 API 승인키
JUSO_CONFIRM_KEY = "DEV100OMP"  # 테스트용 키 또는 사용자님의 실제 발급 키 입력

def fix_zipcode(val):
    if pd.isna(val) or val is None:
        return ""
    val_str = str(val).split('.')[0].strip()
    return val_str.zfill(5) if val_str else ""

SPECIAL_EXCEPTIONS = {
    "불로동 268-2": "인천광역시 검단구 금정로 12 (불로동, 신검단중앙역풍경채어바니티)",
    "남산타운": "서울특별시 중구 다산로 32 (신당동, 남산타운)",
    "율곡동 415": "경상북도 김천시 혁신8로 66",
}

def remove_illegal_chars(val):
    if isinstance(val, str):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', val)
    return val

def remove_stray_province(addr_str):
    if not addr_str:
        return addr_str
    tokens = str(addr_str).split()
    new_tokens = []
    seen_prov = False
    provinces = ['강원특별자치도', '강원도', '서울특별시', '부산광역시', '대구광역시', '인천광역시', '광주광역시', '대전광역시', '울산광역시', '세종특별자치시', '경기도', '충청북도', '충청남도', '전라북도', '전라남도', '경상북도', '경상남도', '제주특별자치도']
    for tok in tokens:
        if tok in provinces:
            if not seen_prov:
                seen_prov = True
                new_tokens.append(tok)
            else:
                continue
        else:
            new_tokens.append(tok)
    return ' '.join(new_tokens)

def preprocess_attached_dong(addr_str):
    if not addr_str:
        return addr_str
    s = str(addr_str).strip()
    s = re.sub(r'([가-힣]+팰리스)([A-Z가-힣\d]동)', r'\1 \2', s)
    s = re.sub(r'([가-힣]+하이츠)([A-Z가-힣\d]동)', r'\1 \2', s)
    s = re.sub(r'([가-힣]+아파트)([A-Z가-힣\d]동)', r'\1 \2', s)
    return s

def remove_duplicate_words(addr_str):
    if not addr_str:
        return addr_str
    s = str(addr_str).strip()
    words = s.split()
    cleaned_words = []
    i = 0
    while i < len(words):
        cleaned_words.append(words[i])
        matched_dup = False
        for length in range(4, 0, -1):
            if i + length * 2 <= len(words):
                seq1 = words[i:i+length]
                seq2 = words[i+length:i+length*2]
                if seq1 == seq2:
                    i += length
                    matched_dup = True
                    break
        if not matched_dup:
            i += 1
    return ' '.join(cleaned_words)

def clean_product_name(val):
    if pd.isna(val) or not str(val).strip():
        return val
    s = str(val).strip()
    
    has_film = False
    if ('풀액정필름2장' in s or '추가액정필름' in s) and '필요없음' not in s and ('풀액정' in s or '액정필름' in s):
        has_film = True

    target_texts = [s]
    if '//' in s: target_texts.insert(0, s.split('//')[-1])
    if '모델선택:' in s: target_texts.insert(0, s.split('모델선택:')[-1])
    if '/' in s: target_texts.insert(0, s.split('/')[-1])

    smartphone_match = None
    for t in target_texts:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', t, re.IGNORECASE)
        if smartphone_match: break

    if not smartphone_match:
        smartphone_match = re.search(r'(S\s*10\s*5G|S\s*10\s*플러스|S\s*20\s*플러스|S\s*20\s*울트라|S\s*22\s*플러스|S\s*22\s*울트라|노트\s*10\s*플러스|노트\s*20\s*울트라|노트\s*20\s*플러스|노트\s*\d+\s*울트라|노트\s*\d+\s*플러스|노트\s*\d+|S\s*\d+(?:\s*플러스|\s*울트라|\s*5G|\s*엣지)?|A\s*\d+|Z\s*(?:플립|폴드)\s*\d*|와이드\s*\d+|N\s*\d+|V\s*50S|V\s*50|V\s*\d+|벨벳|윙)', s, re.IGNORECASE)

    if smartphone_match:
        raw_model = smartphone_match.group(1).replace(" ", "").upper()
        if raw_model == 'S105G': target_model = 'S10 5G'
        elif raw_model == 'S10플러스': target_model = 'S10플러스'
        elif raw_model == 'S20플러스': target_model = 'S20플러스'
        elif raw_model == 'S22플러스': target_model = 'S22플러스'
        elif raw_model == '노트10플러스': target_model = '노트10플러스'
        elif raw_model in ['N960', 'SM-N960']: target_model = '노트9'
        elif raw_model == 'V50S': target_model = 'V50S'
        elif raw_model == 'V50': target_model = 'V50'
        elif raw_model == '노트20울트라': target_model = '노트20울트라'
        elif raw_model == '와이드5': target_model = 'A22'
        elif raw_model == '와이드6': target_model = 'A13'
        else: target_model = raw_model

        qty = 1
        if '☞' in s:
            parts = s.split('☞')
            after = parts[1].strip() if len(parts) > 1 else ""
            qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
            if qty_match: qty = int(qty_match.group(1))
        
        res = f"{target_model}[{qty}]" if qty > 1 else target_model
        if has_film: res += " +필름"
        if '택배' in s or '택배배송' in s: res += " 택배"
        return res
    
    if '케이스' not in s:
        res = s
        if has_film: res += " +필름"
        return res
    
    qty = 1
    if '☞' in s:
        parts = s.split('☞')
        base = parts[0].strip()
        after = parts[1].strip() if len(parts) > 1 else ""
        qty_match = re.match(r'^(\d+)\s*(EA|개)?', after, re.IGNORECASE)
        if qty_match: qty = int(qty_match.group(1))
    else:
        base = s
        
    if '모델선택:' in base:
        parts = base.split('모델선택:')
        target = parts[-1].strip()
        if '/' in target: target = target.split('/')[0].strip()
        if ',' in target: target = target.split(',')[0].strip()
    else:
        sub_parts = re.split(r'[/]{1,2}', base)
        target = sub_parts[-1].strip() if len(sub_parts) > 1 else base
        
    target = re.sub(r'\([^)]*\)', '', target).strip()
    target = re.sub(r'케이스', '', target)
    target = re.sub(r'풀액정', '', target)
    target = re.sub(r'2장', '', target)
    target = re.sub(r'클리어', '', target)
    target = re.sub(r'투명', '', target)
    target = re.sub(r'갤럭시', '', target)
    target = ' '.join(target.split())
    
    if target == '와이드5': target = 'A22'
    elif target == '와이드6': target = 'A13'
    
    result_str = f"{target}[{qty}]" if qty > 1 else target
    if '택배' in s or '택배배송' in s: result_str = f"{result_str} 택배"
    if has_film: result_str = f"{result_str} +필름"
        
    return result_str

# --- [ 행정안전부 주소정보누리집 API 연동 주소 변환 엔진 ] ---
def juso_gov_converter(keyword, original_zip=""):
    if not keyword or pd.isna(keyword):
        return keyword
        
    kw_str = str(keyword).strip()
    kw_str = preprocess_attached_dong(kw_str)
    
    for target_key, override_addr in SPECIAL_EXCEPTIONS.items():
        if target_key in kw_str:
            return override_addr
            
    # 1. 괄호 원본 분리
    parens = re.findall(r'\([^)]+\)', kw_str)
    for p in parens:
        kw_str = kw_str.replace(p, "").strip()
    paren_str = parens[-1] if parens else ""

    # 2. 상세 동·호수 및 배송 메모 추출
    tail_detail = ""
    m_detail = re.search(r'(\d+동\s*\d+호.*|[A-Z가-힣]\s*동\s*\d+호.*|\d+\s*호.*)', kw_str)
    if m_detail:
        tail_detail = m_detail.group(1).strip()
        kw_str = kw_str[:m_detail.start()].strip()

    # 3. 구주소 패턴 검사 및 번지수 기준 API 쿼리 추출
    has_jibun_pattern = bool(re.search(r'([가-힣]+(시|도))\s+([가-힣]+(시|군|구)?(\s+[가-힣]+구)?)\s+([가-힣]+(동|읍|면|리))\s+\d+', kw_str))
    
    base_road_addr = ""
    tail_part = ""
    
    if has_jibun_pattern:
        m_jibun = re.search(r'([가-힣]+동\s*\d+(?:-\d+)?)', kw_str)
        if m_jibun:
            end_pos = m_jibun.end()
            api_query_part = kw_str[:end_pos].strip()
            tail_part = kw_str[end_pos:].strip()
        else:
            api_query_part = kw_str
            tail_part = ""

        base_road_addr = api_query_part
        
        # 행정안전부 주소정보누리집 오픈 API 호출
        url = "https://www.juso.go.kr/addrlink/addrLinkApi.do"
        params = {
            "confmKey": JUSO_CONFIRM_KEY,
            "currentPage": 1,
            "countPerPage": 5,
            "keyword": api_query_part,
            "resultType": "json"
        }
        try:
            response = requests.get(url, params=params, timeout=5)
            if response.status_code == 200:
                res = response.json()
                juso_list = res.get('results', {}).get('juso', [])
                if juso_list:
                    # 도로명 주소(roadAddr) 사용
                    road_addr = juso_list[0].get('roadAddr', '')
                    if road_addr:
                        base_road_addr = road_addr
        except Exception:
            pass
            
        if not base_road_addr:
            base_road_addr = api_query_part
    else:
        base_road_addr = kw_str

    # 4. 최종 조립 순서: [기본 주소] -> [괄호(원본)] -> [상세 동·호수 및 메모]
    parts = [base_road_addr]
    if tail_part and tail_part not in base_road_addr:
        parts.append(tail_part)
    if paren_str:
        parts.append(paren_str)
    if tail_detail:
        parts.append(tail_detail)

    final_str = ' '.join(parts)
    final_str = remove_duplicate_words(final_str)
    final_str = remove_stray_province(final_str)
    final_str = re.sub(r'\(\s*\)', '', final_str)
    return ' '.join(final_str.split())

# --- [ Streamlit 웹 UI 실행 부문 ] ---
st.set_page_config(page_title="주소정보누리집 API 주소 변환기", page_icon="🏛️", layout="centered")

st.title("🏛️ 스마트샵 주소 변환 & 엑셀 정제 (주소정보누리집 API)")
st.write("행정안전부 도로명주소 API를 연동한 최종 버전입니다.")

uploaded_file = st.file_uploader("변환할 엑셀 파일(.xlsx, .xls, .xlsb)을 업로드하세요", type=["xlsx", "xls", "xlsb"])

if uploaded_file is not None:
    try:
        df = pd.read_excel(uploaded_file)
    except Exception:
        df = pd.read_excel(uploaded_file, engine='xlrd')

    st.write("### 📄 업로드 데이터 미리보기 (상위 5건)")
    st.dataframe(df.head())

    if 'excel_data' not in st.session_state:
        st.session_state.excel_data = None

    if st.button("🚀 주소 변환 및 서식 적용 시작"):
        with st.spinner("행정안전부 API로 주소를 변환하는 중입니다..."):
            if '우편번호' in df.columns:
                df['우편번호'] = df['우편번호'].apply(fix_zipcode)

            if '선택정보' in df.columns:
                df['선택정보'] = df['선택정보'].apply(clean_product_name)

            if '배송지' in df.columns:
                progress_bar = st.progress(0)
                total = len(df)
                
                converted_addrs = []
                for i, row in df.iterrows():
                    orig_zip = str(row['우편번호']).zfill(5) if '우편번호' in df.columns and not pd.isna(row['우편번호']) else ""
                    res = juso_gov_converter(row['배송지'], original_zip=orig_zip)
                    converted_addrs.append(res)
                    progress_bar.progress((i + 1) / total)
                    
                df['배송지'] = converted_addrs
            else:
                st.warning("[주의] '배송지' 컬럼을 찾을 수 없습니다.")

            for col in df.columns:
                df[col] = df[col].apply(remove_illegal_chars)

            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False)
                worksheet = writer.sheets['Sheet1']
                
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    for cell in col:
                        val = str(cell.value or '')
                        length = sum(2 if ord(char) > 127 else 1 for char in val)
                        if length > max_len:
                            max_len = length
                    adjusted_width = max(max_len + 4, 12)
                    worksheet.column_dimensions[col_letter].width = min(adjusted_width, 80)

                font_size_8 = Font(size=8)
                for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=len(df.columns)):
                    worksheet.row_dimensions[row[0].row].height = 18
                    for cell in row:
                        cell.font = font_size_8
                        
                if '우편번호' in df.columns:
                    zip_col_idx = df.columns.get_loc('우편번호') + 1
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(row=row, column=zip_col_idx).number_format = '@'

            st.session_state.excel_data = output.getvalue()
        st.success("🎉 행정안전부 API 주소 변환 및 정리가 완료되었습니다!")

    if st.session_state.excel_data is not None:
        st.download_button(
            label="📥 변환된 엑셀 파일 다운로드",
            data=st.session_state.excel_data,
            file_name="우편배송_행정안전부API_최종정리.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
